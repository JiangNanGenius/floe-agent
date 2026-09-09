#!/usr/bin/env python3
"""Independently inspect synthetic native editor outputs; fail on fidelity loss.

This checks the named fixture properties only. Visual fidelity, original-file
CAS, Microsoft Office interoperability and physical devices need separate tests.
"""
import argparse
import hashlib
import json
from pathlib import Path
from docx import Document
from openpyxl import load_workbook


def inspect(original, saved):
    checks = []
    def check(name, expected, actual):
        checks.append({'name': name, 'expected': expected, 'actual': actual, 'passed': actual == expected})
    if original.suffix == '.docx':
        before, after = Document(original), Document(saved)
        check('paragraph count', len(before.paragraphs), len(after.paragraphs))
        check('title preserved', before.paragraphs[0].text, after.paragraphs[0].text)
        check('precise pasted marker retained', True, 'EDITED_NATIVE_WORD ' in after.paragraphs[1].text)
        check('existing paragraph retained', True, after.paragraphs[1].text.endswith(before.paragraphs[1].text))
        check('blue text preserved', before.paragraphs[2].text, after.paragraphs[2].text)
        check('blue font sizes', [18.0], sorted({run.font.size.pt for run in after.paragraphs[2].runs if run.text and run.font.size}))
        check('blue font colors', ['165DBE'], sorted({str(run.font.color.rgb) for run in after.paragraphs[2].runs if run.text}))
        check('table contents', [[c.text for c in r.cells] for r in before.tables[0].rows], [[c.text for c in r.cells] for r in after.tables[0].rows])
        check('header', before.sections[0].header.paragraphs[0].text, after.sections[0].header.paragraphs[0].text)
        for field in ['page_width', 'page_height', 'top_margin', 'bottom_margin', 'left_margin', 'right_margin']:
            check(field, int(getattr(before.sections[0], field)), int(getattr(after.sections[0], field)))
    elif original.suffix == '.xlsx':
        before, after = load_workbook(original), load_workbook(saved)
        values = load_workbook(saved, data_only=True)
        check('sheet order', before.sheetnames, after.sheetnames)
        first, sheet = before['Quarterly'], after['Quarterly']
        check('edited B2 value', 36, sheet['B2'].value)
        check('C2 formula', '=B2*2', sheet['C2'].value)
        check('C2 recalculated cache', 72, values['Quarterly']['C2'].value)
        for cell in ['A1', 'B1', 'C1', 'A2', 'A3', 'B3', 'C3']:
            check(f'{cell} value preserved', first[cell].value, sheet[cell].value)
        for cell in ['A1', 'B2', 'B3', 'C2', 'C3']:
            check(f'{cell} number format', first[cell].number_format, sheet[cell].number_format)
        check('header size', first['A1'].font.sz, sheet['A1'].font.sz)
        check('header weight', first['A1'].font.bold, sheet['A1'].font.bold)
        check('header fill RGB', first['A1'].fill.fgColor.rgb[-6:], sheet['A1'].fill.fgColor.rgb[-6:])
        check('column A width', first.column_dimensions['A'].width, sheet.column_dimensions['A'].width)
        check('chart count', len(first._charts), len(sheet._charts))
        if first._charts and sheet._charts:
            check('chart anchor', [4, 1, 0, 0], [sheet._charts[0].anchor._from.col, sheet._charts[0].anchor._from.row,
                                                sheet._charts[0].anchor._from.colOff, sheet._charts[0].anchor._from.rowOff])
        check('notes sheet content', before['Notes']['A1'].value, after['Notes']['A1'].value)
    else:
        raise ValueError('No verified fixture contract for this extension yet')
    return {'format': original.suffix, 'sourceSHA256': hashlib.sha256(original.read_bytes()).hexdigest(),
            'savedSHA256': hashlib.sha256(saved.read_bytes()).hexdigest(),
            'checks': checks, 'allNamedChecksPassed': all(item['passed'] for item in checks),
            'completeVisualFidelityPassed': False, 'originalFileWritebackPassed': False,
            'microsoftOfficeReopenPassed': False, 'physicalDevicePassed': False}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('original', type=Path)
    parser.add_argument('saved', type=Path)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    result = inspect(args.original, args.saved)
    args.output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + '\n')
    print(json.dumps({'allNamedChecksPassed': result['allNamedChecksPassed'],
                      'failures': [item for item in result['checks'] if not item['passed']]}, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result['allNamedChecksPassed'] else 1)

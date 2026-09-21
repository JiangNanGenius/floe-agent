#!/usr/bin/env python3
"""Create synthetic native editor fixtures using established OOXML libraries.

Run with the bundled Python runtime containing python-docx, openpyxl and
python-pptx. Output is private test input, never a user document.
"""
from pathlib import Path
import argparse
from docx import Document
from docx.shared import Pt, RGBColor as WordColor
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt as SlidePt


def create(output):
    output.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading('Floe Office 原生编辑验证', 0)
    document.add_paragraph('ROUNDTRIP_WORD — edit this paragraph using the native editor.')
    run = document.add_paragraph().add_run('蓝色文字 · 18 pt · format must survive')
    run.font.size = Pt(18)
    run.font.color.rgb = WordColor(0x16, 0x5D, 0xBE)
    table = document.add_table(rows=3, cols=2)
    table.style = 'Table Grid'
    for row, values in zip(table.rows, [('项目', '数值'), ('Alpha', '12'), ('Beta', '24')]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.sections[0].header.paragraphs[0].text = 'Synthetic header — preserve position'
    document.save(output / 'fixture.docx')

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Quarterly'
    for values in [('Item', 'Value', 'Double'), ('Alpha', 12, '=B2*2'), ('Beta', 24, '=B3*2')]:
        sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF', size=14)
        cell.fill = PatternFill('solid', fgColor='165DBE')
    sheet.column_dimensions['A'].width = 22
    sheet.column_dimensions['B'].width = 14
    sheet.column_dimensions['C'].width = 14
    chart = BarChart()
    chart.title = 'Synthetic values'
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
    chart.width = 13
    chart.height = 8
    sheet.add_chart(chart, 'E2')
    workbook.create_sheet('Notes')['A1'] = 'ROUNDTRIP_EXCEL — preserve sheets and formulas'
    workbook.save(output / 'fixture.xlsx')

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    title = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(12), Inches(.7))
    title.text_frame.paragraphs[0].text = 'ROUNDTRIP_POWERPOINT · 原生编辑验证'
    title.text_frame.paragraphs[0].font.size = SlidePt(26)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.8), Inches(2), Inches(3), Inches(2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x16, 0x5D, 0xBE)
    box.text = 'Move or resize this object'
    data = CategoryChartData()
    data.categories = ['Alpha', 'Beta']
    data.add_series('Values', (12, 24))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.8), Inches(7), Inches(4.5), data)
    second = deck.slides.add_slide(deck.slide_layouts[6])
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(2)).text = 'Second slide — preserve ordering'
    deck.save(output / 'fixture.pptx')
    create_sample_deck(output / 'sample-deck.pptx')


def create_sample_deck(path):
    """The deterministic presentation used by the visible-render qualification.

    `verify_pptx_deck_semantics.py` asserts every string, geometry and chart
    value below from the saved package alone, so a rendered-tile receipt can be
    tied to a known deck. Keep this function and that verifier in sync.
    """
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])          # slide 1
    title = slide.shapes.add_textbox(Inches(.7), Inches(.5), Inches(12), Inches(.8))
    title.text_frame.paragraphs[0].text = 'FLOE SAMPLE DECK — 幻灯片一'
    title.text_frame.paragraphs[0].font.size = SlidePt(28)
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.8), Inches(2), Inches(3), Inches(2))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(0x16, 0x5D, 0xBE)
    box.text = 'SAMPLE_BOX'
    data = CategoryChartData()
    data.categories = ['Alpha', 'Beta']
    data.add_series('Values', (12, 24))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(5), Inches(1.8), Inches(7), Inches(4.5), data)
    second = deck.slides.add_slide(deck.slide_layouts[6])          # slide 2
    second.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(2)).text = 'SAMPLE_SECOND_SLIDE'
    third = deck.slides.add_slide(deck.slide_layouts[6])           # slide 3
    oval = third.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2), Inches(2), Inches(3), Inches(3))
    oval.fill.solid()
    oval.fill.fore_color.rgb = RGBColor(0x00, 0xA6, 0x50)
    oval.text = 'SAMPLE_THIRD_SLIDE'
    deck.save(path)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('output', type=Path)
    create(parser.parse_args().output)

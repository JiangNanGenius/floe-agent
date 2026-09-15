"""Real XML, XPath, XSLT and Office round trips in the iOS testbed."""
import io
import json
import sys
from lxml import etree
from docx import Document
from pptx import Presentation

assert sys.platform == "ios"
root = etree.fromstring('<root><text>中文 English</text></root>'.encode())
assert root.xpath('string(text)') == '中文 English'
transform = etree.XSLT(etree.XML(b'''<xsl:stylesheet version="1.0" xmlns:xsl="http://www.w3.org/1999/XSL/Transform"><xsl:template match="/"><out><xsl:value-of select="root/text"/></out></xsl:template></xsl:stylesheet>'''))
assert str(transform(root)).find('中文 English') >= 0
document = Document()
document.add_paragraph('Floe 中文 Office round trip')
word = io.BytesIO(); document.save(word); word.seek(0)
assert Document(word).paragraphs[0].text == 'Floe 中文 Office round trip'
deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[0])
slide.shapes.title.text = 'Floe 中文 slides'
slides = io.BytesIO(); deck.save(slides); slides.seek(0)
assert Presentation(slides).slides[0].shapes.title.text == 'Floe 中文 slides'
print(json.dumps({'lxml': etree.LXML_VERSION, 'libxml2': etree.LIBXML_VERSION,
                  'libxslt': etree.LIBXSLT_VERSION, 'docx_pptx_round_trip': 'passed'}))
